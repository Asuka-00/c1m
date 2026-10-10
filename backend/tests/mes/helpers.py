import io
from typing import Any

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlmodel import Session

from app.core.config import settings
from tests.api.routes.test_config import _owner

API = f"{settings.API_V1_STR}/mes"


class Mes:
    """Tiny client for /mes/<entity> as the owner of a fresh tenant."""

    def __init__(self, client: TestClient, db: Session) -> None:
        self.c = client
        self.h, self.tid, self.slug = _owner(client, db)

    def url(self, e: str, *p: str) -> str:
        return "/".join([API, e.replace("_", "-"), *p])

    def post(self, e: str, body: dict[str, Any], status: int = 200) -> dict[str, Any]:
        r = self.c.post(self.url(e), headers=self.h, json=body)
        assert r.status_code == status, r.text
        return r.json()  # type: ignore[no-any-return]

    def patch(
        self, e: str, rid: str, body: dict[str, Any], status: int = 200
    ) -> dict[str, Any]:
        r = self.c.patch(self.url(e, rid), headers=self.h, json=body)
        assert r.status_code == status, r.text
        return r.json()  # type: ignore[no-any-return]

    def list(self, e: str, **params: Any) -> dict[str, Any]:
        r = self.c.get(self.url(e), headers=self.h, params=params)
        assert r.status_code == 200, r.text
        return r.json()  # type: ignore[no-any-return]

    def delete(self, e: str, rid: str, status: int = 200) -> None:
        r = self.c.delete(self.url(e, rid), headers=self.h)
        assert r.status_code == status, r.text

    def action(
        self, e: str, rid: str, act: str, body: Any = None, status: int = 200
    ) -> dict[str, Any]:
        r = self.c.post(self.url(e, rid, act), headers=self.h, json=body)
        assert r.status_code == status, r.text
        return r.json()  # type: ignore[no-any-return]

    def export(self, e: str) -> list[list[Any]]:
        r = self.c.get(self.url(e, "export"), headers=self.h)
        assert r.status_code == 200, r.text
        ws = load_workbook(io.BytesIO(r.content)).worksheets[0]
        return [list(row) for row in ws.iter_rows(values_only=True)]

    def import_rows(self, e: str, rows: list[list[Any]], status: int = 200) -> Any:
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        for row in rows:
            ws.append(row)
        buf = io.BytesIO()
        wb.save(buf)
        r = self.c.post(
            self.url(e, "import"),
            headers=self.h,
            files={"file": ("x.xlsx", buf.getvalue(), "application/octet-stream")},
        )
        assert r.status_code == status, r.text
        return r.json()
