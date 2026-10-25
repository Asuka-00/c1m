"""按实体规格生成导入模板 xlsx（app/resources/templates/mes/<实体>.xlsx）。

    uv run python scripts/export_import_templates.py [--check]

仅在内容变化时重写文件，并删除已不存在实体的模板。``--check`` 在模板过期时返回 1（测试同样会检查）。
"""

import io
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def normalized(sheets: dict[str, Any]) -> dict[str, list[list[Any]]]:
    """工作表行去掉行尾空单元格（xlsx 读回时不保留）。"""
    out: dict[str, list[list[Any]]] = {}
    for name, rows in sheets.items():
        out[name] = []
        for r in rows:
            r = list(r)
            while r and r[-1] is None:
                r.pop()
            out[name].append(r)
    return out


def read_sheets(path: Path) -> dict[str, list[list[Any]]]:
    """读取 xlsx 的全部工作表单元格值。"""
    wb = load_workbook(io.BytesIO(path.read_bytes()), read_only=True)
    return normalized(
        {ws.title: ws.iter_rows(values_only=True) for ws in wb.worksheets}
    )


def expected() -> dict[Path, Any]:
    """应存在的模板：路径 → 实体规格（不含只读与内部实体）。"""
    import app.main  # noqa: F401  （导入全部模块的规格）
    from app.core import registry
    from app.core.spec.engine import ENTITIES
    from app.core.spec.excel import template_path

    registry.load()
    return {
        template_path(e): e
        for e in ENTITIES.values()
        if not e.readonly and not e.internal
    }


def stale() -> list[Path]:
    """缺失、过期或多余的模板文件。"""
    from app.core.spec.excel import TEMPLATE_DIR, template_sheets

    want = expected()
    bad = [
        p
        for p, e in want.items()
        if not p.exists() or read_sheets(p) != normalized(template_sheets(e))
    ]
    bad += [p for p in (TEMPLATE_DIR / "mes").glob("*.xlsx") if p not in want]
    return sorted(bad)


def main() -> None:
    from app.core.spec.excel import template_xlsx

    bad = stale()
    if "--check" in sys.argv:
        if bad:
            sys.exit(
                f"stale import templates: {', '.join(p.name for p in bad[:10])} "
                "(run scripts/export_import_templates.py)"
            )
        return
    want = expected()
    for p in bad:
        if p in want:
            p.parent.mkdir(parents=True, exist_ok=True)
            p.write_bytes(template_xlsx(want[p]))
        else:
            p.unlink()
    print(f"updated {len(bad)} of {len(want)} templates")  # noqa: T201


if __name__ == "__main__":
    main()
