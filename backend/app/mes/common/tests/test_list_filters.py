"""通用列表筛选：?id= 指定记录、日期 / 时间范围（_from / _to）、时间字段默认降序。"""

from typing import Any

from fastapi.testclient import TestClient
from sqlmodel import Session

from app.core.config import settings
from app.testing.mes import API, Mes


def test_id_range_and_default_desc(client: TestClient, db: Session) -> None:
    m = Mes(client, db)
    a = m.post("site", {"code": "S-A", "name": "a"})
    b = m.post("site", {"code": "S-B", "name": "b"})
    r = m.list("site", id=a["id"])
    assert [x["code"] for x in r["data"]] == ["S-A"]
    r = m.list("site", id=f"{a['id']},{b['id']}")
    assert r["count"] == 2
    assert client.get(m.url("site"), headers=m.h, params={"id": "x"}).status_code == 422
    # 时间范围：created_at（日志）与 datetime 字段
    assert m.list("audit_log", created_at_from="2000-01-01T00:00:00Z")["count"] >= 2
    assert m.list("audit_log", created_at_to="2000-01-01T00:00:00Z")["count"] == 0
    eq = m.post("equipment", {"code": "EQ-R", "name": "eq"})
    rc = m.post("reason_code", {"code": "DT", "name": "d", "reason_type": "downtime"})
    for d in ("2026-01-01", "2026-03-01", "2026-02-01"):
        m.post(
            "downtime",
            {
                "equipment_id": eq["id"],
                "start_at": f"{d}T08:00:00+08:00",
                "reason_id": rc["id"],
            },
        )
    # 默认按开始时间降序（最新在前）
    rows = m.list("downtime")["data"]
    assert [x["start_at"][:7] for x in rows] == ["2026-03", "2026-02", "2026-01"]
    r = m.list(
        "downtime",
        start_at_from="2026-01-15T00:00:00+08:00",
        start_at_to="2026-02-28T23:59:59+08:00",
    )
    assert [x["start_at"][:7] for x in r["data"]] == ["2026-02"]
    bad = client.get(m.url("downtime"), headers=m.h, params={"start_at_from": "nope"})
    assert bad.status_code == 422


def test_export_honours_filters_and_distinct(client: TestClient, db: Session) -> None:
    import io

    from openpyxl import load_workbook

    m = Mes(client, db)
    for c in ("EX-A", "EX-B", "EY-C"):
        m.post("site", {"code": c, "name": c})
    r = client.get(m.url("site", "export"), headers=m.h, params={"q": "EX-"})
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws is not None
    codes = sorted(str(x[0].value) for x in ws.iter_rows(min_row=2) if x[0].value)
    assert [c for c in codes if c.startswith("E")] == ["EX-A", "EX-B"]
    # 文本筛选去重取值：仅 filter=True 的 str 字段
    d = client.get(m.url("audit_log", "distinct", "username"), headers=m.h)
    assert d.status_code == 200 and d.json()["data"]
    bad = client.get(m.url("site", "distinct", "name"), headers=m.h)
    assert bad.status_code == 422


def test_meta_conditional_scope_fields(client: TestClient, db: Session) -> None:
    m = Mes(client, db)
    meta = client.get(f"{API}/_meta", headers=m.h).json()
    e = next(x for x in meta["entities"] if x["code"] == "check_binding")
    f = {x["name"]: x for x in e["fields"]}
    assert f["equipment"]["only_when_required"] is True
    assert f["trigger"]["only_when_required"] is False
    smr = next(x for x in meta["entities"] if x["code"] == "split_merge_rule")
    assert not any(x["filter"] for x in smr["fields"] if x["type"] == "bool")


def test_export_datetime_fields_utc_z(client: TestClient, db: Session) -> None:
    """D46 P2：xlsx 里 datetime 是按租户时区换算的 Excel 日期单元格，列名带时区；
    接口 JSON 仍是 UTC Z。导出文件可原样导回（列名时区后缀被去掉）。"""
    import io
    from datetime import datetime

    from openpyxl import load_workbook

    m = Mes(client, db)
    eq = m.post("equipment", {"code": "EQ-XD", "name": "eq"})
    rc = m.post("reason_code", {"code": "DT-X", "name": "d", "reason_type": "downtime"})
    m.post(
        "downtime",
        {
            "equipment_id": eq["id"],
            "start_at": "2026-04-01T08:00:00+08:00",
            "reason_id": rc["id"],
        },
    )
    r = client.get(m.url("downtime", "export"), headers=m.h)
    assert r.status_code == 200, r.text
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws is not None
    rows = list(ws.iter_rows(values_only=True))
    # 没有厂区：默认 Asia/Taipei
    i = rows[0].index("start_at (Asia/Taipei)")
    assert datetime(2026, 4, 1, 8, 0) in [x[i] for x in rows[1:]]
    assert ws.cell(row=2, column=i + 1).number_format == "yyyy-mm-dd hh:mm:ss"
    assert "equipment" in rows[0]  # 非时间列名不变
    js = m.list("downtime")["data"]
    assert js[0]["start_at"] == "2026-04-01T00:00:00Z"
    # 租户时区取第一个启用厂区
    m.post("site", {"code": "A-UTC", "name": "s", "timezone": "UTC"})
    r = client.get(m.url("downtime", "export"), headers=m.h)
    rows = list(load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True))  # type: ignore[union-attr]
    i = rows[0].index("start_at (UTC)")
    assert datetime(2026, 4, 1, 0, 0) in [x[i] for x in rows[1:]]


def test_export_import_datetime_roundtrip(client: TestClient, db: Session) -> None:
    """D46 P2：导出的 xlsx（列名带时区、日期单元格）可原样导回，时间不偏移。"""
    m = Mes(client, db)
    eq = m.post("equipment", {"code": "EQ-RT", "name": "eq"})
    rc = m.post(
        "reason_code", {"code": "DT-RT", "name": "d", "reason_type": "downtime"}
    )
    m.post(
        "downtime",
        {
            "equipment_id": eq["id"],
            "start_at": "2026-04-01T08:30:00+08:00",
            "reason_id": rc["id"],
        },
    )
    rows = m.export("downtime")
    assert "start_at (Asia/Taipei)" in rows[0]
    m.import_rows("downtime", rows)
    got = [x["start_at"] for x in m.list("downtime")["data"]]
    assert got == ["2026-04-01T00:30:00Z", "2026-04-01T00:30:00Z"], got


def test_range_boundaries_d77(client: TestClient, db: Session) -> None:
    """D77：datetime 的 _to 半开（恰在边界的记录不含）；date 的 _to 闭区间（含当天）。"""
    m = Mes(client, db)
    eq = m.post("equipment", {"code": "EQ-B", "name": "eq"})
    rc = m.post("reason_code", {"code": "DT", "name": "d", "reason_type": "downtime"})
    m.post(
        "downtime",
        {
            "equipment_id": eq["id"],
            "start_at": "2026-03-10T00:00:00+08:00",
            "reason_id": rc["id"],
        },
    )

    def dt(**kw: str) -> int:
        return int(m.list("downtime", **kw)["count"])

    assert dt(start_at_to="2026-03-10T00:00:00+08:00") == 0  # 边界不含
    assert dt(start_at_to="2026-03-09T16:00:00Z") == 0  # 同一时刻（UTC 写法）
    assert dt(start_at_to="2026-03-10T00:00:00.000001+08:00") == 1
    assert dt(start_at_from="2026-03-10T00:00:00+08:00") == 1  # 下限含
    assert (
        dt(
            start_at_from="2026-03-09T16:00:00Z",
            start_at_to="2026-03-10T00:00:01+08:00",
        )
        == 1
    )
    # date 字段：qualified_until = 3/10，_to=3/10 含当天，_to=3/9 不含
    op = m.post("operation", {"code": "OP-B", "name": "op"})
    m.post(
        "equipment_capability",
        {
            "equipment_id": eq["id"],
            "operation_id": op["id"],
            "qualified_until": "2026-03-10",
        },
    )

    def cap(**kw: str) -> int:
        return int(m.list("equipment_capability", **kw)["count"])

    assert cap(qualified_until_to="2026-03-10") == 1
    assert cap(qualified_until_to="2026-03-09") == 0
    assert cap(qualified_until_from="2026-03-10", qualified_until_to="2026-03-10") == 1
    assert cap(qualified_until_from="2026-03-11") == 0


def test_range_d77_supplement(client: TestClient, db: Session) -> None:
    """D77 补充一：datetime 必须带偏移；from > to 统一 range_inverted；from == to 边界。"""
    m = Mes(client, db)
    eq = m.post("equipment", {"code": "EQ-S", "name": "eq"})
    rc = m.post("reason_code", {"code": "DT", "name": "d", "reason_type": "downtime"})
    m.post(
        "downtime",
        {
            "equipment_id": eq["id"],
            "start_at": "2026-03-10T00:00:00+08:00",
            "reason_id": rc["id"],
        },
    )
    op = m.post("operation", {"code": "OP-S", "name": "op"})
    m.post(
        "equipment_capability",
        {
            "equipment_id": eq["id"],
            "operation_id": op["id"],
            "qualified_until": "2026-03-10",
        },
    )

    def get(ent: str, **kw: str) -> Any:
        return client.get(m.url(ent), headers=m.h, params=kw)

    # datetime：只给日期 / 无偏移 → timezone_required
    for name in ("start_at_from", "start_at_to"):
        for raw in ("2026-03-10", "2026-03-10T00:00:00"):
            r = get("downtime", **{name: raw})
            assert r.status_code == 422, (name, raw)
            d = r.json()["detail"][0]
            assert (d["type"], d["loc"]) == ("timezone_required", ["query", name])
    r = get("audit_log", created_at_from="2026-03-10")
    assert r.json()["detail"][0]["type"] == "timezone_required"
    # date 字段只给日期照常
    assert (
        get("equipment_capability", qualified_until_to="2026-03-10").json()["count"]
        == 1
    )
    # from > to：date 与 datetime 都是 value_error + range_inverted
    for ent, kw in (
        (
            "downtime",
            {
                "start_at_from": "2026-03-10T00:00:01+08:00",
                "start_at_to": "2026-03-10T00:00:00+08:00",
            },
        ),
        (
            "equipment_capability",
            {"qualified_until_from": "2026-03-11", "qualified_until_to": "2026-03-10"},
        ),
    ):
        r = get(ent, **kw)
        assert r.status_code == 422, r.text
        d = r.json()["detail"][0]
        assert d["type"] == "value_error" and d["ctx"] == {"reason": "range_inverted"}
        assert d["loc"] == ["query", next(iter(kw))]
    # from == to：datetime 半开 → 200 空；date 闭区间 → 当天
    same = "2026-03-10T00:00:00+08:00"
    r = get("downtime", start_at_from=same, start_at_to=same)
    assert r.status_code == 200 and r.json()["count"] == 0
    r = get(
        "equipment_capability",
        qualified_until_from="2026-03-10",
        qualified_until_to="2026-03-10",
    )
    assert r.status_code == 200 and r.json()["count"] == 1


def test_date_field_from_offset_datetime_uses_site_tz(
    client: TestClient, db: Session
) -> None:
    """D75：date 字段收到带时区的时间，先换算到站点时区（Asia/Taipei）再取日期。"""
    m = Mes(client, db)
    eq = m.post("equipment", {"code": "EQ-T", "name": "eq"})
    op = m.post("operation", {"code": "OP-T", "name": "op"})
    c = m.post(
        "equipment_capability",
        {
            "equipment_id": eq["id"],
            "operation_id": op["id"],
            "qualified_until": "2026-03-09T20:00:00Z",  # 台北 03-10 04:00
        },
    )
    assert c["qualified_until"] == "2026-03-10"


def test_import_error_loc_points_to_cell(client: TestClient, db: Session) -> None:
    """Excel 导入错误 loc 为 ["file", 行号, 列名]，并给 column / cell。"""
    m = Mes(client, db)
    m.post("equipment", {"code": "EQ-I", "name": "eq"})
    rows = m.export("equipment_capability")
    header = list(rows[0])
    row: list[Any] = [None] * len(header)
    row[header.index("equipment")] = "EQ-I"
    row[header.index("operation")] = "NO-SUCH-OP"
    err = m.import_rows("equipment_capability", [header, row], 422)
    e0 = err["detail"]["errors"][0]
    i = header.index("operation")
    assert e0["row"] == 2 and e0["loc"] == ["file", 2, "operation"], err
    assert e0["cell"] == f"{chr(65 + i)}2", err


def test_date_filter_strict_pattern_d77_supp2(client: TestClient, db: Session) -> None:
    """D77 补充二：date 筛选只收严格 YYYY-MM-DD，带时间 → value_error/pattern。"""
    m = Mes(client, db)
    for raw in (
        "2026-03-10T00:00:00+08:00",
        "2026-03-10T00:00:00",
        "2026-03-10 00:00",
        "2026-3-10",
        "20260310",
        "2026/10/01",
        "2026-02-30",
    ):
        r = client.get(
            m.url("equipment_capability"),
            headers=m.h,
            params={"qualified_until_to": raw},
        )
        assert r.status_code == 422, raw
        d = r.json()["detail"][0]
        assert d["type"] == "value_error", raw
        assert d["ctx"] == {"reason": "pattern"}
        assert d["loc"] == ["query", "qualified_until_to"]
    r = client.get(
        m.url("equipment_capability"),
        headers=m.h,
        params={"qualified_until_from": "2026-03-10"},
    )
    assert r.status_code == 200


def test_report_window_range_inverted_d77_supp2(
    client: TestClient, db: Session
) -> None:
    """D77 补充二：output/oee/yield/oee-trend 报表 from>to → range_inverted。"""
    m = Mes(client, db)
    for path in ("output", "oee", "yield", "oee-trend"):
        r = client.get(
            f"{settings.API_V1_STR}/mes/exec/reports/{path}",
            headers=m.h,
            params={"date_from": "2026-03-11", "date_to": "2026-03-10"},
        )
        assert r.status_code == 422, (path, r.text)
        d = r.json()["detail"][0]
        assert (d["type"], d["loc"], d["ctx"]) == (
            "value_error",
            ["query", "date_from"],
            {"reason": "range_inverted"},
        ), path


def test_date_equality_and_datetime_garbage_pattern(
    client: TestClient, db: Session
) -> None:
    """date 字段等值筛选同样 pattern；datetime 无法解析 → pattern，只给日期 → timezone_required。"""
    m = Mes(client, db)

    def err(ent: str, **kw: str) -> Any:
        r = client.get(m.url(ent), headers=m.h, params=kw)
        assert r.status_code == 422, (ent, kw, r.text)
        return r.json()["detail"][0]

    d = err("equipment_capability", qualified_until="2026/10/01")
    assert (d["type"], d["loc"], d["ctx"]) == (
        "value_error",
        ["query", "qualified_until"],
        {"reason": "pattern"},
    )
    for name in ("start_at_from", "start_at_to"):
        d = err("downtime", **{name: "2026/10/01"})
        assert (d["type"], d["ctx"]) == ("value_error", {"reason": "pattern"})
        assert err("downtime", **{name: "2026-10-01"})["type"] == "timezone_required"
    assert (
        client.get(
            m.url("equipment_capability"),
            headers=m.h,
            params={"qualified_until": "2026-10-01"},
        ).status_code
        == 200
    )


def test_custom_routes_dates_d77(client: TestClient, db: Session) -> None:
    """自定义路由：shift-logs（date）pattern；holds / alarms / maps（datetime）D77 补充一。"""
    m = Mes(client, db)
    api = f"{settings.API_V1_STR}/mes"

    def err(path: str, **kw: str) -> Any:
        r = client.get(f"{api}/{path}", headers=m.h, params=kw)
        assert r.status_code == 422, (path, kw, r.text)
        return r.json()["detail"][0]

    for raw in ("2026/10/01", "2026-10-01T00:00:00"):
        d = err("exec/shift-logs", date_from=raw)
        assert (d["type"], d["loc"], d["ctx"]) == (
            "value_error",
            ["query", "date_from"],
            {"reason": "pattern"},
        )
    d = err("exec/shift-logs", date_from="2026-10-02", date_to="2026-10-01")
    assert d["ctx"] == {"reason": "range_inverted"}
    for path, f, t in (
        ("exec/holds", "date_from", "date_to"),
        ("eqp/alarms", "from", "to"),
        ("map/maps", "imported_from", "imported_to"),
    ):
        d = err(path, **{f: "2026-10-01"})
        assert (d["type"], d["loc"]) == ("timezone_required", ["query", f]), path
        d = err(path, **{t: "2026-10-01T00:00:00"})
        assert (d["type"], d["loc"]) == ("timezone_required", ["query", t]), path
        d = err(path, **{f: "2026/10/01"})
        assert d["ctx"] == {"reason": "pattern"}, path
        d = err(path, **{f: "2026-10-02T00:00:00Z", t: "2026-10-01T00:00:00Z"})
        assert (d["loc"], d["ctx"]) == (["query", f], {"reason": "range_inverted"})
        ok = client.get(
            f"{api}/{path}",
            headers=m.h,
            params={f: "2026-10-01T00:00:00+08:00", t: "2026-10-02T00:00:00Z"},
        )
        assert ok.status_code == 200, (path, ok.text)


def test_report_window_strict_date_d77_supp3(client: TestClient, db: Session) -> None:
    """D77 补充三：报表 date_from/date_to 只收严格 YYYY-MM-DD；格式先于 range_inverted。"""
    m = Mes(client, db)
    for path in ("output", "oee", "yield", "oee-trend"):
        url = f"{settings.API_V1_STR}/mes/exec/reports/{path}"
        for name in ("date_from", "date_to"):
            for raw in (
                "2026-03-10T00:00:00",
                "2026-03-10T00:00:00+08:00",
                "2026/03/10",
            ):
                r = client.get(url, headers=m.h, params={name: raw})
                assert r.status_code == 422, (path, name, raw, r.text)
                d = r.json()["detail"][0]
                assert (d["type"], d["loc"], d["ctx"]) == (
                    "value_error",
                    ["query", name],
                    {"reason": "pattern"},
                ), (path, name, raw)
        # 格式错误先于 range_inverted（from 合法但晚于 to，to 带时间）
        r = client.get(
            url,
            headers=m.h,
            params={"date_from": "2026-03-11", "date_to": "2026-03-10T00:00:00"},
        )
        d = r.json()["detail"][0]
        assert (d["loc"], d["ctx"]) == (["query", "date_to"], {"reason": "pattern"})
        r = client.get(
            url,
            headers=m.h,
            params={"date_from": "2026-03-01", "date_to": "2026-03-10"},
        )
        assert r.status_code == 200, (path, r.text)
