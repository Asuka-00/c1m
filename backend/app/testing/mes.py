import io
import uuid
from typing import Any

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlmodel import Session

from app.core.config import settings
from app.platform.config.tests.test_config import _owner

API = f"{settings.API_V1_STR}/mes"


class Mes:
    """Tiny client for /mes/<entity> as the owner of a fresh tenant."""

    def __init__(self, client: TestClient, db: Session) -> None:
        self.c = client
        self.h, self.tid, self.slug = _owner(client, db)
        self._fx: dict[str, str] = {}

    # Fixture records for required context refs, so tests only spell out what
    # they are about. A key given explicitly (even None) is never filled.
    FIXTURES: dict[str, tuple[str, dict[str, Any]]] = {
        "site": ("site", {"code": "FX-SITE", "name": "fx"}),
        "area": ("area", {"code": "FX-AREA", "name": "fx", "site": "@site"}),
        "workcenter": ("workcenter", {"code": "FX-WC", "name": "fx", "area": "@area"}),
        "uom": ("uom", {"code": "FX-EA", "name": "fx", "dimension": "count"}),
        "tech_node": ("tech_node", {"code": "FX-N", "name": "fx"}),
        "family": (
            "product_family",
            {"code": "FX-FAM", "name": "fx", "tech_node": "@tech_node"},
        ),
        "product": ("product", {"code": "FX-PROD", "name": "fx"}),
        "equipment_type": (
            "equipment_type",
            {"code": "FX-ET", "name": "fx", "category": "process"},
        ),
        "operation": ("operation", {"code": "FX-OP", "name": "fx"}),
        "default_priority": (
            "lot_priority",
            {"code": "FX-PRI", "name": "fx", "level": 50},
        ),
        "lot_id_rule": (
            "code_rule",
            {"code": "FX-LOT", "name": "fx", "target": "lot", "pattern": "FX{SEQ}"},
        ),
        "wafer_id_rule": (
            "code_rule",
            {
                "code": "FX-WFR",
                "name": "fx",
                "target": "wafer",
                "pattern": "{LOT}.{SLOT}",
            },
        ),
        "lot_type": ("lot_type", {"code": "FX-LT", "name": "fx"}),
        "reason": (
            "reason_code",
            {"code": "FX-HOLD", "name": "fx", "reason_type": "hold"},
        ),
    }
    CONTEXT: dict[str, list[str]] = {
        "product_family": ["tech_node"],
        "product": ["uom", "family", "tech_node"],
        "reticle": ["product"],
        "equipment": ["equipment_type", "workcenter"],
        "equipment_group": ["equipment_type"],
        "recipe": ["equipment_type"],
        "operation": ["area"],
        "dc_plan": ["operation"],
        "sampling_rule": ["operation"],
        "lot_type": ["default_priority", "lot_id_rule", "wafer_id_rule"],
        "split_merge_rule": ["lot_type"],
        "param_spec": ["uom"],
        "measurement_item": ["uom"],
        "qtime_rule": ["reason"],
    }

    def fx(self, key: str) -> str:
        if key not in self._fx:
            ent, body = self.FIXTURES[key]
            body = {
                k: self.fx(v[1:]) if isinstance(v, str) and v.startswith("@") else v
                for k, v in body.items()
            }
            self._fx[key] = self.post(ent, body)["id"]
        return self._fx[key]

    def _fill(self, e: str, body: dict[str, Any]) -> dict[str, Any]:
        skip = (
            (e == "equipment" and body.get("unit_kind", "main") != "main")
            or (
                e == "measurement_item"
                and body.get("meas_type", "numeric") != "numeric"
            )
            or (
                e == "qtime_rule"
                and body.get("action", "hold") not in ("hold", "scrap")
            )
        )
        out = dict(body)
        keys = [] if skip else self.CONTEXT.get(e, [])
        if e == "product" and body.get("product_type") in ("monitor", "dummy"):
            keys = ["uom"]
        for k in keys:
            if k not in body and f"{k}_id" not in body:
                out[f"{k}_id"] = self.fx(k)
        return out

    def hk(self, e: str) -> dict[str, str]:
        """受控写实体（D81 KEYED）自动带新 Idempotency-Key。"""
        from app.mes.common.keyed_write import KEYED

        if e not in KEYED:
            return self.h
        return {**self.h, "Idempotency-Key": str(uuid.uuid4())}

    def qual_extra(self, comment: str = "renew") -> dict[str, Any]:
        """资质停用/延期必带的 reason_id / comment / approved_by（D74）；审批人为另一成员。"""
        if "_qual_extra" not in self._fx:
            rc = self.post(
                "reason_code",
                {"code": "OQ-RENEW", "name": "renew", "reason_type": "operator_qual"},
            )
            name = "appr_" + uuid.uuid4().hex[:8]
            r = self.c.post(
                f"{settings.API_V1_STR}/tenants/{self.tid}/members/new",
                headers=self.h,
                json={
                    "username": name,
                    "full_name": "approver",
                    "password": uuid.uuid4().hex,
                    "roles": ["member"],
                },
            )
            assert r.status_code == 200, r.text
            self._fx["_qual_extra"] = rc["id"]
            self._fx["_qual_appr"] = r.json()["username"]
        return {
            "reason_id": self._fx["_qual_extra"],
            "comment": comment,
            "approved_by": self._fx["_qual_appr"],
        }

    def url(self, e: str, *p: str) -> str:
        return "/".join([API, e.replace("_", "-"), *p])

    def post(self, e: str, body: dict[str, Any], status: int = 200) -> dict[str, Any]:
        r = self.c.post(self.url(e), headers=self.hk(e), json=self._fill(e, body))
        assert r.status_code == status, r.text
        return r.json()  # type: ignore[no-any-return]

    def patch(
        self, e: str, rid: str, body: dict[str, Any], status: int = 200
    ) -> dict[str, Any]:
        r = self.c.patch(self.url(e, rid), headers=self.hk(e), json=body)
        assert r.status_code == status, r.text
        return r.json()  # type: ignore[no-any-return]

    def list(self, e: str, **params: Any) -> dict[str, Any]:
        r = self.c.get(self.url(e), headers=self.h, params=params)
        assert r.status_code == 200, r.text
        return r.json()  # type: ignore[no-any-return]

    def delete(self, e: str, rid: str, status: int = 200) -> None:
        r = self.c.delete(self.url(e, rid), headers=self.hk(e))
        assert r.status_code == status, r.text

    def action(
        self, e: str, rid: str, act: str, body: Any = None, status: int = 200
    ) -> dict[str, Any]:
        r = self.c.post(self.url(e, rid, act), headers=self.h, json=body)
        assert r.status_code == status, r.text
        return r.json()  # type: ignore[no-any-return]

    def export(self, e: str) -> list[list[Any]]:
        r = self.c.get(self.url(e, "export"), headers=self.h)
        assert r.status_code == 200, r.text
        ws = load_workbook(io.BytesIO(r.content)).worksheets[0]
        return [list(row) for row in ws.iter_rows(values_only=True)]

    def import_rows(self, e: str, rows: list[list[Any]], status: int = 200) -> Any:
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        for row in rows:
            ws.append(row)
        buf = io.BytesIO()
        wb.save(buf)
        r = self.c.post(
            self.url(e, "import"),
            headers=self.h,
            files={"file": ("x.xlsx", buf.getvalue(), "application/octet-stream")},
        )
        assert r.status_code == status, r.text
        return r.json()
