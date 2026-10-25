"""按实体导出 / 导入 xlsx（openpyxl）。

工作表布局：第 1 行为列键（字段名，版本化数据加 ``version``/``status``，自定义属性为
``attr:<code>``）。引用以目标编码读写（``CODE`` = 激活或最新版本，``CODE@2`` = 指定版本），
无编码的目标使用 id。导入按业务唯一键新增或更新，全部成功或全部回滚。

导入模板（``template_sheets``）由 ``scripts/export_import_templates.py`` 写入
``app/resources/templates/mes/<实体>.xlsx`` 并纳入版本库；测试校验模板与规格一致。
"""

import io
import json
import uuid
from pathlib import Path
from typing import Any

import sqlalchemy as sa
from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from sqlmodel import Session

from app.core import error_codes as ec
from app.core.spec.base import conn, table
from app.core.spec.engine import ENTITIES, ENUMS, Entity, F, label_of
from app.core.spec.errors import CoerceError, coded
from app.core.spec.query import all_rows
from app.core.spec.serialize import to_json
from app.core.spec.validate import attr_defs
from app.core.spec.write import create_row, update_row

MAX_ROWS = 5000
TEMPLATE_DIR = Path(__file__).resolve().parents[2] / "resources" / "templates"
README = "README"
_FORMATS = {
    "int": "integer",
    "float": "number",
    "bool": "TRUE / FALSE",
    "date": "YYYY-MM-DD",
    "time": "HH:MM:SS",
    "datetime": "YYYY-MM-DDTHH:MM:SS",
    "json": "JSON",
}


def _ref_code(
    session: Session, tenant_id: uuid.UUID, target: Entity, ids: set[Any]
) -> dict[Any, str]:
    if not ids:
        return {}
    t = table(target)
    if not target.has_code:
        return {i: str(i) for i in ids}
    cols = [t.c.id, t.c.code] + ([t.c.version] if target.versioned else [])
    out = {}
    for r in (
        conn(session)
        .execute(sa.select(*cols).where(t.c.tenant_id == tenant_id, t.c.id.in_(ids)))
        .mappings()
    ):
        out[r["id"]] = f"{r['code']}@{r['version']}" if target.versioned else r["code"]
    return out


def columns(session: Session, tenant_id: uuid.UUID, e: Entity) -> list[str]:
    """导出/导入的列键：字段名、版本化实体的 version/status、自定义属性 ``attr:<code>``。"""
    cols = [f.name for f in e.all_fields]
    if e.versioned:
        cols += ["version", "status"]
    cols += [f"attr:{d['code']}" for d in attr_defs(session, tenant_id, e)]
    return cols


def export_xlsx(
    session: Session,
    tenant_id: uuid.UUID,
    e: Entity,
    *,
    q: str | None = None,
    filters: dict[str, str] | None = None,
) -> bytes:
    """导出实体行为 xlsx（按当前关键字 / 筛选；无筛选时为全部行）。

    Returns:
        xlsx 文件字节。
    """
    rows = all_rows(session, e, tenant_id, MAX_ROWS, q=q, filters=filters)
    cols = columns(session, tenant_id, e)
    refs = {
        f.name: _ref_code(
            session,
            tenant_id,
            ENTITIES[f.ref or ""],
            {r[f.column] for r in rows if r[f.column]},
        )
        for f in e.all_fields
        if f.type == "ref"
    }
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = e.code[:31]
    ws.append(cols)
    for r in rows:
        line: list[Any] = []
        for c in cols:
            if c.startswith("attr:"):
                line.append((r.get("attrs") or {}).get(c[5:]))
                continue
            f = e.field(c)
            if f and f.type == "ref":
                line.append(refs[f.name].get(r[f.column]))
            elif f and f.type == "json":
                line.append(
                    json.dumps(r[f.column]) if r[f.column] is not None else None
                )
            else:
                v = r[f.column if f else c]
                line.append(to_json(v) if f and f.type in ("date", "time") else v)
        ws.append(line)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def template_path(e: Entity) -> Path:
    """实体导入模板文件路径。"""
    return TEMPLATE_DIR / "mes" / f"{e.code}.xlsx"


def _format(f: F) -> str:
    if f.type == "enum":
        return " / ".join(v for v, _ in ENUMS[f.enum or ""])
    if f.type == "ref":
        target = ENTITIES[f.ref or ""]
        if not target.has_code:
            return f"{target.code} id"
        return f"{target.code} code" + (
            " (CODE or CODE@version)" if target.versioned else ""
        )
    if f.type == "str":
        return f"text (max {f.max_length})"
    return _FORMATS.get(f.type, "text")


def _required(f: F) -> str | None:
    if f.required_when:
        cond = "; ".join(
            f"{k} = {'any' if '*' in v else ' / '.join(map(str, v))}"
            for k, v in f.required_when.items()
        )
        return f"when {cond}"
    if not f.required:
        return None
    return "auto" if f.auto or f.default is not None or f.default_from else "yes"


def template_sheets(e: Entity) -> dict[str, list[list[Any]]]:
    """导入模板内容：首个工作表为导入列键（表头），``README`` 为逐列说明。

    不含系统字段与租户自定义属性（可另加 ``attr:<code>`` 列）。

    Returns:
        工作表名 → 行（单元格值）。
    """
    fields = [f for f in e.all_fields if not f.system]
    head = [f.name for f in fields] + (["version"] if e.versioned else [])
    doc: list[list[Any]] = [["column", "label", "format", "required", "help"]]
    doc += [[f.name, label_of(f), _format(f), _required(f), f.help] for f in fields]
    if e.versioned:
        doc.append(["version", "Version", "integer", "yes", None])
    doc += [
        [None] * 5,
        [
            f"Rows are matched by {', '.join(e.natural_key)}: existing rows are updated, others created."
        ],
        [
            f"Only the first sheet is imported; all rows succeed or nothing is saved (max {MAX_ROWS})."
        ],
        [
            "Leave unused columns blank or delete them; add attr:<code> columns for custom attributes."
        ],
    ]
    return {e.code[:31]: [head], README: doc}


def template_xlsx(e: Entity) -> bytes:
    """按 ``template_sheets`` 生成导入模板 xlsx。"""
    wb = Workbook()
    assert wb.active is not None
    wb.remove(wb.active)
    for name, rows in template_sheets(e).items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
        for c in ws[1]:
            c.font = Font(bold=True)
        for col in ws.iter_cols(max_row=len(rows)):
            width = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(width, 10), 40)
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _resolve(session: Session, tenant_id: uuid.UUID, target: Entity, raw: str) -> Any:
    t = table(target)
    c = conn(session)
    if not target.has_code:
        try:
            return uuid.UUID(raw)
        except ValueError:
            return None
    code, _, ver = raw.partition("@")
    stmt = sa.select(t.c.id).where(t.c.tenant_id == tenant_id, t.c.code == code.strip())
    if target.versioned:
        if ver:
            stmt = stmt.where(t.c.version == int(ver))
        else:
            stmt = stmt.order_by(
                sa.case((t.c.status == "active", 0), else_=1), t.c.version.desc()
            )
    return c.execute(stmt.limit(1)).scalar()


def _row_errors(n: int, d: Any) -> list[dict[str, Any]]:
    """把一行的 HTTPException 详情展开为行错误（保留错误码与参数供前端翻译）。"""
    if isinstance(d, list):
        out = []
        for x in d:
            item: dict[str, Any] = {"row": n, "msg": x.get("msg", "")}
            if x.get("code"):
                item["code"], item["params"] = x["code"], x.get("params") or {}
            loc = x.get("loc") or []
            if loc:
                item["field"] = str(loc[-1])
            out.append(item)
        return out
    if isinstance(d, dict):
        item = {"row": n, "msg": str(d.get("message", ""))}
        if d.get("code"):
            item["code"], item["params"] = d["code"], d.get("params") or {}
        return [item]
    return [{"row": n, "msg": str(d)}]


def import_xlsx(
    session: Session, tenant_id: uuid.UUID, e: Entity, data: bytes, user: Any
) -> dict[str, int]:
    """从 xlsx 导入（按业务唯一键新增或更新）。

    Args:
        session: 数据库会话。
        tenant_id: 租户 id。
        e: 实体规格。
        data: xlsx 文件字节。
        user: 操作人。

    Returns:
        ``{"created": n, "updated": m}``。

    Raises:
        HTTPException: 任一行校验失败（422，消息带行号）。
    """
    if e.readonly:
        raise coded(405, ec.READ_ONLY_ENTITY, "Read-only entity")
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as ex:  # noqa: BLE001 - any parse error is a bad upload
        raise coded(400, ec.NOT_A_VALID_XLSX_FILE, "Not a valid xlsx file") from ex
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(it, ())]
    if not header or not any(header):
        raise coded(400, ec.EMPTY_SHEET, "Empty sheet")
    known = set(columns(session, tenant_id, e)) | {""}
    unknown = [h for h in header if h not in known]
    if unknown:
        raise coded(
            422,
            ec.IMPORT_UNKNOWN_COLUMNS,
            f"Unknown columns: {', '.join(unknown)}",
            columns=", ".join(unknown),
        )
    errors: list[dict[str, Any]] = []
    created = updated = 0
    t = table(e)
    nested = session.begin_nested()
    for n, values in enumerate(it, start=2):
        if n > MAX_ROWS + 1:
            errors.append(
                {
                    "row": n,
                    "msg": f"too many rows (max {MAX_ROWS})",
                    "code": ec.IMPORT_TOO_MANY_ROWS,
                    "params": {"max": MAX_ROWS},
                }
            )
            break
        raw = dict(zip(header, values, strict=False))
        if all(v in (None, "") for v in raw.values()):
            continue
        payload: dict[str, Any] = {}
        attrs: dict[str, Any] = {}
        try:
            for k, v in raw.items():
                if not k or k == "status":
                    continue
                if k.startswith("attr:"):
                    if v not in (None, ""):
                        attrs[k[5:]] = v
                    continue
                f = e.field(k)
                if f and f.type == "ref" and v not in (None, ""):
                    rid = _resolve(session, tenant_id, ENTITIES[f.ref or ""], str(v))
                    if rid is None:
                        raise CoerceError(
                            ec.REF_VALUE_NOT_FOUND,
                            f"{k}: '{v}' not found",
                            field=k,
                            value=str(v),
                        )
                    payload[f.column] = rid
                elif f and f.type == "json" and isinstance(v, str) and v.strip():
                    payload[k] = json.loads(v)
                else:
                    payload[k] = v
            if attrs:
                payload["attrs"] = attrs
            key_cols = [
                (fld.column if (fld := e.field(k)) else k) for k in e.natural_key
            ]
            key_vals = [
                payload.get(c, payload.get(c.removesuffix("_id"))) for c in key_cols
            ]
            existing = None
            if all(v not in (None, "") for v in key_vals):
                stmt = sa.select(t.c.id).where(t.c.tenant_id == tenant_id)
                for c, v in zip(key_cols, key_vals, strict=True):
                    column: Any = t.c[c]
                    stmt = stmt.where(column == (int(str(v)) if c == "version" else v))
                existing = conn(session).execute(stmt).scalar()
            if existing:
                update_row(
                    session,
                    tenant_id,
                    e,
                    existing,
                    payload,
                    user,
                    commit=False,
                    action="import",
                )
                updated += 1
            else:
                create_row(session, tenant_id, e, payload, user, commit=False)
                created += 1
        except HTTPException as ex:
            errors.extend(_row_errors(n, ex.detail))
        except CoerceError as ex:
            errors.append(
                {"row": n, "msg": str(ex), "code": ex.code, "params": ex.params}
            )
        except (ValueError, TypeError) as ex:
            errors.append({"row": n, "msg": str(ex)})
    if errors:
        nested.rollback()
        session.rollback()
        raise HTTPException(
            status_code=422, detail={"message": "Import failed", "errors": errors[:50]}
        )
    nested.commit()
    session.commit()
    return {"created": created, "updated": updated}
